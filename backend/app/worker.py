"""Single worker with committed claims, renewable leases and fenced publication."""
import io
import logging
import re
import threading
import time
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import or_, select, update

from . import models as m
from .api.reports import snapshot_access
from .common import DomainError, audit, canonical, digest, lock_identity, lock_project, require
from .config import settings
from .db import SessionLocal
from .security import membership
from .services import generate

LEASE_SECONDS = 60
HEARTBEAT_SECONDS = 20
MAX_ATTEMPTS = 3


def safe_cell(value):
    value = '' if value is None else str(value)
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else value


def stable_zip(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, raw in sorted(entries):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, raw)
    return output.getvalue()


def build_export(manifest):
    """CPU/file generation has no open database transaction. Every cell is text."""
    book = Workbook(); book.remove(book.active)
    fixed = datetime.fromisoformat(manifest['as_of']).replace(tzinfo=None)
    book.properties.created = fixed; book.properties.modified = fixed
    book.properties.creator = '财务资料与整改工作台'
    def sheet(name, headings, rows):
        ws = book.create_sheet(name); ws.append(headings)
        for values in rows:
            ws.append(['' if v is None else str(v) for v in values])
        for row in ws:
            for cell in row:
                cell.data_type = 's'  # Explicit string, including = + - @; never a formula.
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='244B45')
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 24
    sheet('快照说明', ['项目', '范围版本', '模板版本', '口径', '截至时间', '用途', 'manifest哈希'], [[manifest['project_id'],
          manifest['scope_version'], manifest['template_version_id'], manifest['metric_version'], manifest['as_of'], manifest['purpose'], digest(manifest)]])
    sheet('资料清单', ['主题', '标题', '主体', '期间', '适用性', '状态', '经办', '复核', '截止日'], [
        [i[k] for k in ('topic_code', 'title', 'org_name', 'period_label', 'applicability', 'state', 'owner_name', 'reviewer_name', 'due')]
        for i in manifest['items']])
    sheet('整改事项', ['事项', '事实', '主体', '等级', '状态', '经办', '验证人', '原截止日', '当前截止日', '延期次数'], [
        [i[k] for k in ('title', 'facts', 'org_id', 'severity', 'state', 'owner_name', 'verifier_name', 'original_due', 'current_due', 'extension_count')]
        for i in manifest['issues']])
    sheet('目录版本', ['目录代号', '主体编号', '期间编号', '版本', '保管部门', '位置代号', '取得状态', 'metadata_hash'], [
        [e['code'], e['org_id'], e['period_id'], e['version']['number'], e['version']['content']['department'],
         e['version']['content']['location_code'], e['version']['content']['acquisition'], e['version']['metadata_hash']]
        for e in manifest['evidence']])
    raw = io.BytesIO(); book.save(raw)
    # openpyxl updates modified time when saving; normalize XML and ZIP metadata.
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source:
        entries = []
        for name in source.namelist():
            content = source.read(name)
            if name == 'docProps/core.xml':
                content = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',
                    lambda match: match[1] + fixed.isoformat().encode() + b'Z' + match[2], content)
            entries.append((name, content))
    workbook = stable_zip(entries)
    return stable_zip([('directory.xlsx', workbook), ('manifest.json', canonical(manifest).encode()),
        ('checksums.json', canonical({'manifest_sha256': digest(manifest), 'directory.xlsx': digest(workbook),
                                     'format': 'directory-xlsx-v1'}).encode())])


def claim_one():
    with SessionLocal.begin() as db:
        instant = m.now()
        job = db.scalar(select(m.Job).where(or_(m.Job.state == 'queued',
            (m.Job.state == 'retry_wait') & (m.Job.next_run_at <= instant),
            (m.Job.state == 'running') & (m.Job.lease_until <= instant)))
            .order_by(m.Job.created_at, m.Job.id).with_for_update(skip_locked=True).limit(1))
        if not job: return None
        if job.attempts >= MAX_ATTEMPTS:
            job.state = 'failed'; job.error_code = 'RETRIES_EXHAUSTED'; job.attempt_token = None
            audit(db, None, job.id, 'job_exhausted', job.project_id, result='failed')
            return ('', '')
        job.state = 'running'; job.attempts += 1; job.attempt_token = m.uid()
        job.heartbeat_at = instant; job.lease_until = instant + timedelta(seconds=LEASE_SECONDS)
        job.error_code = None
        return job.id, job.attempt_token


def renew(job_id, token):
    with SessionLocal.begin() as db:
        instant = m.now()
        return db.execute(update(m.Job).where(m.Job.id == job_id, m.Job.attempt_token == token,
            m.Job.state == 'running', m.Job.lease_until > instant).values(
                heartbeat_at=instant, lease_until=instant + timedelta(seconds=LEASE_SECONDS))).rowcount == 1


def heartbeat(job_id, token, stop):
    while not stop.wait(HEARTBEAT_SECONDS):
        try:
            if not renew(job_id, token): return
        except Exception:
            logging.error('job_heartbeat_failed job_id=%s', job_id)
            return


def authorized_job(db, job_id, token):
    lock_identity(db)
    job = db.get(m.Job, job_id); lock_project(db, job.project_id)
    actor = db.get(m.User, job.actor_id)
    require(actor and actor.active, 'ACTOR_DISABLED', '任务发起人已停用')
    mem = membership(db, actor, job.project_id, ['cfo', 'pmo'])
    require(set(job.payload.get('org_ids', [])) <= set(mem.org_ids), 'JOB_SCOPE_REVOKED', '任务范围授权已撤销')
    require(job.state == 'running' and job.attempt_token == token and job.lease_until > m.now(),
            'LEASE_LOST', '任务租约已失效', 409)
    return job, actor


def fence(db, job_id, token, result):
    # The compare-and-swap locks the job row only at publication. A reclaimed
    # attempt cannot commit business writes even if its old process is still alive.
    count = db.execute(update(m.Job).where(m.Job.id == job_id, m.Job.state == 'running',
        m.Job.attempt_token == token, m.Job.lease_until > m.now()).values(
            state='succeeded', result=result, lease_until=None, attempt_token=None)).rowcount
    require(count == 1, 'LEASE_LOST', '任务租约已失效', 409)


def export_context(db, actor, approval_id):
    approval = db.get(m.Approval, approval_id)
    require(approval and approval.state == 'approved', 'EXPORT_UNAPPROVED', '导出尚未获批')
    require(approval.created_at + timedelta(hours=settings.export_ttl_hours) > m.now(), 'EXPORT_EXPIRED', '导出许可已失效')
    snap = snapshot_access(db, actor, db.get(m.Snapshot, approval.snapshot_id))
    require(approval.payload['manifest_hash'] == snap.manifest_hash == digest(snap.manifest), 'MANIFEST_MISMATCH', '快照校验失败')
    return approval, snap


def execute(job_id, token):
    with SessionLocal.begin() as db:
        job, actor = authorized_job(db, job_id, token)
        if job.kind == 'generate':
            result = generate(db, actor, job.project_id, job.payload)
            fence(db, job_id, token, result); audit(db, actor, job_id, 'job_completed', job.project_id)
            return
        require(job.kind == 'export', 'JOB_KIND_UNKNOWN', '任务类型不支持', 422)
        approval, snap = export_context(db, actor, job.payload['approval_id'])
        manifest = snap.manifest
    blob = build_export(manifest)
    root = settings.export_storage_root; root.mkdir(parents=True, exist_ok=True, mode=0o700)
    temp = root / f'{job_id}-{token}.tmp'; path = root / f'{job_id}-{token}.zip'
    try:
        temp.write_bytes(blob); temp.chmod(0o600)
        with SessionLocal.begin() as db:
            job, actor = authorized_job(db, job_id, token)
            approval, snap = export_context(db, actor, job.payload['approval_id'])
            old = db.scalar(select(m.Export).where(m.Export.approval_id == approval.id))
            if old:
                fence(db, job_id, token, {'export_id': old.id})
                return
            row = m.Export(approval_id=approval.id, snapshot_id=snap.id, path=str(path.resolve()), artifact_hash=digest(blob),
                           expires_at=approval.created_at + timedelta(hours=settings.export_ttl_hours))
            db.add(row); db.flush()
            fence(db, job_id, token, {'export_id': row.id})
            audit(db, actor, row.id, 'export_generated', job.project_id)
            temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)
        # A final archive left by an interrupted commit is removed by orphan cleanup.


def fail_attempt(job_id, token, exc):
    with SessionLocal.begin() as db:
        job = db.scalar(select(m.Job).where(m.Job.id == job_id).with_for_update())
        if job.state != 'running' or job.attempt_token != token: return
        terminal = isinstance(exc, DomainError) or job.attempts >= MAX_ATTEMPTS
        job.state = 'failed' if terminal else 'retry_wait'
        job.error_code = exc.code if isinstance(exc, DomainError) else 'EXECUTION_FAILED'
        job.next_run_at = m.now() + timedelta(seconds=5 if job.attempts == 1 else 15)
        job.lease_until = None; job.attempt_token = None
        audit(db, None, job.id, 'job_failed' if terminal else 'job_retry_scheduled', job.project_id, result='failed')


def process_one():
    claim = claim_one()
    if claim is None: return False
    job_id, token = claim
    if not job_id: return True
    stop = threading.Event(); thread = threading.Thread(target=heartbeat, args=(job_id, token, stop), daemon=True)
    thread.start()
    try:
        execute(job_id, token)
    except Exception as exc:
        fail_attempt(job_id, token, exc)
        logging.error('job_attempt_failed job_id=%s error_type=%s', job_id, type(exc).__name__)
    finally:
        stop.set(); thread.join(timeout=2)
    return True


def cleanup_exports():
    """Retention is idempotent; audit failure rolls back before files are unlinked."""
    root = settings.export_storage_root.resolve()
    with SessionLocal.begin() as db:
        rows = db.scalars(select(m.Export).where(m.Export.expires_at <= m.now(), m.Export.deleted_at.is_(None))
                          .with_for_update(skip_locked=True)).all()
        paths = []
        for row in rows:
            path = Path(row.path).resolve()
            require(path.is_relative_to(root), 'EXPORT_PATH_INVALID', '导出位置无效', 500)
            paths.append(path); row.deleted_at = m.now()
            audit(db, None, row.id, 'export_expired')
    # Also revisit already-marked paths, allowing recovery after a crash between
    # metadata commit and unlink. No directory trees or unknown names are deleted.
    with SessionLocal() as db:
        paths = [Path(p).resolve() for p in db.scalars(select(m.Export.path).where(m.Export.deleted_at.is_not(None)))]
        referenced = set(db.scalars(select(m.Export.path)))
        active_attempts = {f'{j.id}-{j.attempt_token}' for j in db.scalars(select(m.Job).where(
            m.Job.state == 'running', m.Job.lease_until > m.now()))}
    for path in paths:
        if path.is_relative_to(root): path.unlink(missing_ok=True)
    if root.exists():
        for path in root.iterdir():
            if re.fullmatch(r'[0-9a-f-]{36}-[0-9a-f-]{36}\.(tmp|zip)', path.name) and not path.is_symlink():
                if path.stem not in active_attempts and str(path.resolve()) not in referenced and time.time() - path.stat().st_mtime > LEASE_SECONDS * 3:
                    path.unlink(missing_ok=True)


if __name__ == '__main__':
    logging.basicConfig(level=logging.INFO)
    last_cleanup = 0
    while True:
        try:
            if time.monotonic() - last_cleanup > 60:
                cleanup_exports(); last_cleanup = time.monotonic()
            if not process_one(): time.sleep(1)
        except Exception as exc:
            logging.error('worker_failure error_type=%s', type(exc).__name__)
            time.sleep(3)
